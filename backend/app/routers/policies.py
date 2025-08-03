from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List
from datetime import date, datetime
from .. import models, schemas
from app.database import SessionLocal
from .users import get_current_user
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from io import BytesIO
import os


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


router = APIRouter(dependencies=[Depends(get_current_user)], prefix="/polizas", tags=["Polizas"])


@router.post("/", response_model=schemas.PolicyOut)
def create_policy(
    data: schemas.PolicyCreate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    existing = (
        db.query(models.Policy)
        .filter(models.Policy.PolicyNum == data.PolicyNum)
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail="Número de póliza ya existe")
    policy = models.Policy(
        DocType=data.DocType,
        PolicyNum=data.PolicyNum,
        InitDate=data.InitDate,
        DueDate=data.DueDate,
        ComiPrcnt=data.ComiPrcnt,
        AscValue=data.AscValue,
        CreateDate=date.today(),
        LastDateMod=date.today(),
        id_slrs=data.id_slrs,
        id_ctms=data.id_ctms,
        id_usrs_create=current_user.id,
        id_usrs_update=current_user.id,
        id_insurance=data.id_insurance,
        id_poliza_rel=data.id_poliza_rel,
        comentario=data.comentario,
        activo=data.activo or "Y",
        aut_noti=data.aut_noti or "N",
    )
    db.add(policy)
    if data.DocType == "R" and data.id_poliza_rel:
        prev_policy = db.query(models.Policy).get(data.id_poliza_rel)
        if prev_policy:
            prev_policy.activo = "N"
    db.commit()
    db.refresh(policy)
    for line in data.lines:
        pl = models.PolicyLine(
            id_policy=policy.id,
            id_itm=line.id_itm,
            LineNum=line.LineNum,
            LineTotal=line.LineTotal,
        )
        db.add(pl)
    db.commit()
    return policy


@router.get("/", response_model=List[schemas.PolicyListOut])
def list_policies(db: Session = Depends(get_db)):
    policies = db.query(models.Policy).all()
    today = date.today()
    result = []
    for p in policies:
        ins = db.query(models.InsuranceCompany).get(p.id_insurance)
        days_overdue = (today - p.DueDate).days if p.DueDate < today else 0
        base = schemas.PolicyOut.from_orm(p).dict()
        base.update(
            {
                "InsuranceName": ins.CompanyName if ins else None,
                "ComiPrcnt": p.ComiPrcnt,
                "DaysOverdue": days_overdue,
                "RelatedPolicyNum": db.query(models.Policy.PolicyNum).filter(models.Policy.id == p.id_poliza_rel).scalar() if p.id_poliza_rel else None,
            }
        )
        result.append(schemas.PolicyListOut(**base))
    return result

@router.get("/export")
def export_policies(db: Session = Depends(get_db)):
    policies = db.query(models.Policy).all()
    today = date.today()

    wb = Workbook()
    ws = wb.active

    logo_path = os.path.join(
        os.path.dirname(__file__),
        "../../..",
        "frontend",
        "public",
        "LoginForm",
        "logo.png",
    )
    if os.path.exists(logo_path):
        img = XLImage(logo_path)
        img.height = 80
        img.width = 160
        ws.add_image(img, "A1")

    ws.merge_cells("C1:J1")
    ws["C1"] = "Listado de pólizas"
    ws["C1"].font = Font(size=16, bold=True)
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("C2:J2")
    ws["C2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws["C2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.append([])
    headers = [
        "ID",
        "Tipo",
        "Número",
        "Fecha inicio",
        "Fecha vencimiento",
        "% Comisión",
        "Valor asegurado",
        "Aseguradora",
        "Días vencido",
        "Activo",
        "Notificación",
    ]
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True)

    for p in policies:
        ins = db.query(models.InsuranceCompany).get(p.id_insurance)
        days_overdue = (today - p.DueDate).days if p.DueDate < today else 0
        ws.append(
            [
                p.id,
                p.DocType,
                p.PolicyNum,
                p.InitDate.isoformat(),
                p.DueDate.isoformat(),
                p.ComiPrcnt,
                p.AscValue,
                ins.CompanyName if ins else "",
                days_overdue,
                p.activo,
                p.aut_noti,
            ]
        )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"polizas_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

@router.put("/aut-noti")
def update_aut_noti(
    payload: schemas.BulkAutNotiUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    for pid in payload.policy_ids:
        policy = db.query(models.Policy).get(pid)
        if policy:
            policy.aut_noti = payload.aut_noti
            policy.LastDateMod = date.today()
            policy.id_usrs_update = current_user.id
    db.commit()
    return {"msg": "Polizas actualizadas"}

@router.put("/activo")
def update_activo(
    payload: schemas.BulkActivoUpdate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    for pid in payload.policy_ids:
        policy = db.query(models.Policy).get(pid)
        if policy:
            policy.activo = payload.activo
            policy.LastDateMod = date.today()
            policy.id_usrs_update = current_user.id
    db.commit()
    return {"msg": "Polizas actualizadas"}


@router.get("/{id}", response_model=schemas.PolicyDetailOut)
def get_policy(id: int, db: Session = Depends(get_db)):
    policy = db.query(models.Policy).get(id)
    if not policy:
        raise HTTPException(status_code=404, detail="Póliza no encontrada")
    lines = db.query(models.PolicyLine).filter_by(id_policy=id).all()
    return {**policy.__dict__, "lines": lines}


@router.put("/{id}", response_model=schemas.PolicyOut)
def update_policy(
    id: int,
    data: schemas.PolicyCreate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    policy = db.query(models.Policy).get(id)
    if not policy:
        raise HTTPException(status_code=404, detail="Póliza no encontrada")
    
    existing = (
        db.query(models.Policy)
        .filter(models.Policy.PolicyNum == data.PolicyNum, models.Policy.id != id)
        .first()
    )
    if existing:
        raise HTTPException(status_code=400, detail="Número de póliza ya existe")

    policy.DocType = data.DocType
    policy.PolicyNum = data.PolicyNum
    policy.InitDate = data.InitDate
    policy.DueDate = data.DueDate
    policy.ComiPrcnt = data.ComiPrcnt
    policy.AscValue = data.AscValue
    policy.LastDateMod = date.today()
    policy.id_slrs = data.id_slrs
    policy.id_ctms = data.id_ctms
    policy.id_usrs_update = current_user.id
    policy.id_insurance = data.id_insurance
    policy.id_poliza_rel = data.id_poliza_rel
    policy.comentario = data.comentario
    policy.activo = data.activo or policy.activo
    policy.aut_noti = data.aut_noti or policy.aut_noti

    if data.DocType == "R" and data.id_poliza_rel:
        prev_policy = db.query(models.Policy).get(data.id_poliza_rel)
        if prev_policy:
            prev_policy.activo = "N"

    db.query(models.PolicyLine).filter_by(id_policy=id).delete()
    for line in data.lines:
        db.add(
            models.PolicyLine(
                id_policy=id,
                id_itm=line.id_itm,
                LineNum=line.LineNum,
                LineTotal=line.LineTotal,
            )
        )
    db.commit()
    db.refresh(policy)
    return policy


@router.delete("/{id}")
def delete_policy(id: int, db: Session = Depends(get_db)):
    policy = db.query(models.Policy).get(id)
    if not policy:
        raise HTTPException(status_code=404, detail="Póliza no encontrada")
    db.query(models.PolicyLine).filter_by(id_policy=id).delete()
    db.delete(policy)
    db.commit()
    return {"msg": "Póliza eliminada"}

