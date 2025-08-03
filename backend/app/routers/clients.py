from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List
from .. import models, schemas
from .users import get_current_user
from app.database import SessionLocal
from geopy.geocoders import Nominatim
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font
from io import BytesIO
from datetime import datetime
import os

geolocator = Nominatim(user_agent="hagesa_api")

def geocode_cliente(db: Session, data: schemas.ClientBase):
    parts = []
    if data.id_parroquia:
        p = db.query(models.Parish).get(data.id_parroquia)
        if p:
            parts.append(p.Description)
    if data.id_ciudad:
        c = db.query(models.City).get(data.id_ciudad)
        if c:
            parts.append(c.Description)
    if data.id_provincia:
        s = db.query(models.State).get(data.id_provincia)
        if s:
            parts.append(s.Description)
    if data.id_pais:
        pais = db.query(models.Country).get(data.id_pais)
        if pais:
            parts.append(pais.Description)
    if not parts:
        return None, None
    address = ", ".join(parts)
    try:
        loc = geolocator.geocode(address)
        if loc:
            return loc.latitude, loc.longitude
    except Exception:
        pass
    return None, None

def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()

router = APIRouter(dependencies=[Depends(get_current_user)], prefix="/clientes", tags=["Clientes"])

@router.post("/", response_model=schemas.ClientOut)
def create_cliente(cliente: schemas.ClientCreate, db: Session = Depends(get_db)):
    data = cliente.dict()
    lat, lon = geocode_cliente(db, cliente)
    data["latitud"] = lat
    data["longitud"] = lon
    nuevo = models.Client(**data)
    db.add(nuevo)
    db.commit()
    db.refresh(nuevo)
    return nuevo

@router.get("/", response_model=List[schemas.ClientOut])
def list_clientes(db: Session = Depends(get_db)):
    return db.query(models.Client).all()

@router.get("/export")
def export_clientes(db: Session = Depends(get_db)):
    clientes = db.query(models.Client).all()

    wb = Workbook()
    ws = wb.active

    logo_path = os.path.join(
        os.path.dirname(__file__),
        "../../..",
        "frontend",
        "public",
        "LoginForm",
        "logo.png",
    )
    if os.path.exists(logo_path):
        img = XLImage(logo_path)
        img.height = 80
        img.width = 160
        ws.add_image(img, "A1")

    ws.merge_cells("C1:H1")
    ws["C1"] = "Listado de clientes"
    ws["C1"].font = Font(size=16, bold=True)
    ws["C1"].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells("C2:H2")
    ws["C2"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    ws["C2"].alignment = Alignment(horizontal="center", vertical="center")

    ws.append([])
    headers = [
        "ID",
        "Nombre",
        "Apellidos",
        "Identificación",
        "Teléfono",
        "Email",
        "Dirección",
    ]
    ws.append(headers)
    for cell in ws[4]:
        cell.font = Font(bold=True)

    for c in clientes:
        ws.append(
            [
                c.id,
                c.nombre,
                c.apellidos or "",
                c.identificacion,
                c.telefono,
                c.email,
                c.direccion,
            ]
        )

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"clientes_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )

@router.get("/{id}", response_model=schemas.ClientOut)
def get_cliente(id: int, db: Session = Depends(get_db)):
    cliente = db.query(models.Client).get(id)
    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    return cliente

@router.put("/{id}", response_model=schemas.ClientOut)
def update_cliente(id: int, data: schemas.ClientCreate, db: Session = Depends(get_db)):
    cliente = db.query(models.Client).get(id)
    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    for key, value in data.dict().items():
        setattr(cliente, key, value)
    lat, lon = geocode_cliente(db, data)
    cliente.latitud = lat
    cliente.longitud = lon
    db.commit()
    db.refresh(cliente)
    return cliente

@router.delete("/{id}")
def delete_cliente(id: int, db: Session = Depends(get_db)):
    cliente = db.query(models.Client).get(id)
    if not cliente:
        raise HTTPException(status_code=404, detail="Cliente no encontrado")
    db.delete(cliente)
    db.commit()
    return {"msg": "Cliente eliminado"}
